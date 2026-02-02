from __future__ import annotations

from pathlib import Path

from docx import Document


def extract_docx_to_txt(docx_path: Path, out_path: Path) -> None:
    doc = Document(str(docx_path))
    lines: list[str] = []

    lines.append("=== PARAGRAPHS ===")
    for i, p in enumerate(doc.paragraphs, 1):
        text = (p.text or "").rstrip()
        if text:
            lines.append(f"P{i}: {text}")

    lines.append("")
    lines.append("=== TABLES ===")
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"-- Table {ti} --")
        for ri, row in enumerate(table.rows, 1):
            cells = [" ".join((c.text or "").split()) for c in row.cells]
            lines.append(f"R{ri}: " + " | ".join(cells))

    out_path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    src = Path(r"c:\Users\yanboyang\Desktop\project\FYP Interim Report.docx")
    out = Path(r"c:\Users\yanboyang\Desktop\project\_FYP_Interim_Report_extracted.txt")
    extract_docx_to_txt(src, out)
    print(f"WROTE: {out}")
